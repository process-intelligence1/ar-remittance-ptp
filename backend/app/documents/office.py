"""DOCX (paragraphs + tables in document order, headers and footers), XLSX (computed values, dates
as ISO), CSV, HTML and plain-text attachments."""
import csv
import datetime as dt
import io

from app.documents.common import ParseResult
from app.documents.tables import find_html_tables, find_text_tables


# --------------------------------------------------------------------------- DOCX
def _docx_table_rows(table) -> list[list[str]]:
    rows = []
    for r in table.rows:
        cells, seen = [], set()
        for c in r.cells:
            if id(c._tc) in seen:  # a merged cell appears once per spanned column: keep it once
                continue
            seen.add(id(c._tc))
            cells.append(" ".join(c.text.split()))
        if any(cells):
            rows.append(cells)
    return rows


def _docx_blocks(container):
    """Paragraphs and tables in the order they appear (python-docx lists them separately).
    `container` is the document body (doc._body), a header/footer, or a table cell."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in container._element.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            yield Paragraph(child, container)
        elif tag == "tbl":
            yield Table(child, container)


def parse_docx(data: bytes) -> ParseResult:
    import docx

    d = docx.Document(io.BytesIO(data))
    parts: list[str] = []
    tables: list[dict] = []

    def take(container, where: str):
        for block in _docx_blocks(container):
            if block.__class__.__name__ == "Paragraph":
                if block.text.strip():
                    parts.append(block.text.strip())
            else:
                rows = _docx_table_rows(block)
                if rows:
                    tables.append({"page": None, "source": "docx", "name": where, "rows": rows})
                    parts.append("\n".join("\t".join(r) for r in rows))

    headers, footers = [], []
    for section in d.sections:
        for part, bucket in ((section.header, headers), (section.footer, footers)):
            if part is not None and not part.is_linked_to_previous:
                bucket.append(part)
    for h in headers:
        take(h, "header")
    take(d._body, "body")  # the <w:body> element, not <w:document>, holds the paragraphs
    for f in footers:
        take(f, "footer")
    return ParseResult("docx", "\n".join(parts), [], tables)


# --------------------------------------------------------------------------- XLSX
def cell_value(v) -> str:
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat() if v.time() == dt.time(0, 0) else v.isoformat(timespec="seconds")
    if isinstance(v, (dt.date, dt.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_xlsx(data: bytes) -> ParseResult:
    from openpyxl import load_workbook

    # data_only=True: the value Excel last calculated, not the formula text
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    parts, tables = [], []
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [cell_value(v) for v in row]
            while cells and cells[-1] == "":
                cells.pop()
            if any(cells):
                rows.append(cells)
        if rows:
            tables.append({"page": None, "source": "xlsx", "name": ws.title, "rows": rows})
            parts.append(f"[{ws.title}]\n" + "\n".join("\t".join(r) for r in rows))
    wb.close()
    return ParseResult("spreadsheet", "\n\n".join(parts), [], tables, page_count=len(wb.worksheets))


# --------------------------------------------------------------------------- CSV / text / HTML
def decode_text(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_csv(data: bytes) -> ParseResult:
    text = decode_text(data)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect) if any(c.strip() for c in r)]
    tables = [{"page": None, "source": "csv", "rows": rows}] if rows else []
    return ParseResult("csv", "\n".join("\t".join(r) for r in rows), [], tables)


def parse_plain(data: bytes) -> ParseResult:
    text = decode_text(data)
    tables = [{"page": None, "source": "text", "rows": t["rows"]} for t in find_text_tables(text)]
    return ParseResult("plain", text.strip(), [], tables)


def parse_html(data: bytes) -> ParseResult:
    from app.gmail.parsing import html_to_text

    html = decode_text(data)
    tables = [{"page": None, "source": "html", "rows": t} for t in find_html_tables(html)]
    return ParseResult("html", html_to_text(html), [], tables)
