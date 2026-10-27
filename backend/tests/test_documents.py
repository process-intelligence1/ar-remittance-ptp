"""M02 document reading: every reader, OCR, deskew, timeout, reuse by SHA-256, body tables."""
import datetime as dt
import io

import cv2
import numpy as np
import pymupdf
import pytest
from fastapi.testclient import TestClient
from PIL import Image
from sqlalchemy import select

from app.config import get_settings
from app.documents import ocr as ocr_mod
from app.documents.body import parse_body
from app.documents.common import Deadline, DocumentTimeout, detect_language, readable_ratio
from app.documents.pdf import page_needs_ocr
from app.documents.router import kind_of, parse_document
from app.documents.service import parse_email
from app.documents.tables import find_html_tables, find_text_tables
from app.gmail.parsing import html_to_text
from app.gmail.poller import store_email
from app.gmail.sample_sender import build_message, load_sample
from app.main import app
from app.models.ar import EmailAttachment

LINES = ["REMITTANCE ADVICE", "Payment reference: TRF48197765", "",
         "Invoice No.          Amount paid", "UK-1043360           9,841.00", "UK-1043671           1,673.51"]


@pytest.fixture(autouse=True)
def tesseract_path():
    ocr_mod._configure_tesseract()


def text_pdf(lines=LINES, pages=1) -> bytes:
    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page()
        for i, line in enumerate(lines):
            page.insert_text((72, 90 + 22 * i), line, fontsize=13, fontname="cour")  # monospace keeps columns
    return doc.tobytes()


def render_png(lines=LINES, dpi=200) -> bytes:
    page = pymupdf.open(stream=text_pdf(lines), filetype="pdf")[0]
    return page.get_pixmap(dpi=dpi, colorspace=pymupdf.csGRAY).tobytes("png")


def scanned_pdf(extra_text_page=False) -> bytes:
    """A PDF whose page is only an image (like a scan); optionally followed by a normal text page."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(page.rect, stream=render_png(dpi=200))
    if extra_text_page:
        doc.insert_pdf(pymupdf.open(stream=text_pdf(), filetype="pdf"))
    return doc.tobytes()


# --------------------------------------------------------------------------- PDF
def test_text_pdf_words_and_tables():
    r = parse_document("advice.pdf", "application/pdf", text_pdf())
    assert (r.method, r.status, r.page_count, r.ocr_confidence) == ("text", "PARSED", 1, None)
    assert "UK-1043360" in r.text and r.language == "en"
    page = r.pages[0]
    assert page["unit"] == "pt" and page["source"] == "text"
    word = next(w for w in page["words"] if w[0] == "UK-1043360")
    assert all(0 <= v <= 1 for v in word[1:5]) and word[1] < word[3] and word[2] < word[4]


def test_scanned_pdf_is_ocrd():
    r = parse_document("scan.pdf", None, scanned_pdf())
    assert r.method == "ocr" and r.pages[0]["source"] == "ocr" and r.pages[0]["reason"] == "scanned"
    assert r.ocr_confidence and r.ocr_confidence > 70
    line = next(ln for ln in r.text.splitlines() if "UK-1043360" in ln)
    assert "9,841.00" in line  # invoice number and amount stay on the same line


def test_mixed_pdf():
    r = parse_document("mixed.pdf", None, scanned_pdf(extra_text_page=True))
    assert r.method == "mixed" and r.page_count == 2
    assert [p["source"] for p in r.pages] == ["ocr", "text"]


def test_unreadable_text_layer_is_ocrd():
    assert page_needs_ocr("short") == "scanned"
    garbage = "  (cid:12)(cid:15)(cid:18) ��" * 5
    assert readable_ratio(garbage) < 0.7 and page_needs_ocr(garbage) == "unreadable text layer"
    assert page_needs_ocr("Invoice UK-1043360 amount 9,841.00 paid on 21 Oct 2026") is None


def test_pdf_page_limit(monkeypatch):
    monkeypatch.setattr(get_settings(), "pdf_max_pages", 2)
    r = parse_document("long.pdf", None, text_pdf(pages=4))
    assert len(r.pages) == 2 and r.page_count == 4 and "first 2 of 4" in r.note


# --------------------------------------------------------------------------- images / OCR
def test_png_ocr_keeps_rows():
    r = parse_document("advice.png", "image/png", render_png())
    assert r.method == "ocr" and r.page_count == 1 and r.ocr_confidence > 70
    assert any("UK-1043671" in ln and "1,673.51" in ln for ln in r.text.splitlines())
    assert any(["UK-1043360", "9,841.00"] == row for t in r.tables for row in t["rows"])


def test_deskew_straightens_a_rotated_scan():
    gray = np.asarray(Image.open(io.BytesIO(render_png(dpi=150))))
    h, w = gray.shape
    rotated = cv2.warpAffine(gray, cv2.getRotationMatrix2D((w / 2, h / 2), 5, 1.0), (w, h),
                             borderValue=255)
    assert abs(abs(ocr_mod.skew_angle(rotated)) - 5) < 1.0
    ok, png = cv2.imencode(".png", rotated)
    r = parse_document("tilted.png", None, png.tobytes())
    assert "UK-1043360" in r.text


def test_large_image_is_scaled_down(monkeypatch):
    monkeypatch.setattr(get_settings(), "ocr_max_side_px", 1000)
    prepared, info = ocr_mod.prepare_image(Image.new("L", (3000, 2000), 255))
    assert max(prepared.shape) <= 1000 and info["scale"] < 1


def test_ocr_timeout():
    with pytest.raises(DocumentTimeout):
        parse_document("advice.png", None, render_png(), timeout_s=0.001)


def test_build_page_joins_japanese_without_spaces():
    data = {"text": ["振込", "番号", "TRF96500401"], "conf": [90, 90, 95], "left": [0, 42, 200],
            "top": [0, 0, 0], "width": [40, 40, 120], "height": [20, 20, 20],
            "block_num": [1, 1, 1], "par_num": [1, 1, 1], "line_num": [1, 1, 1]}
    page = ocr_mod.build_page(data, 400, 100, "jpn+eng")
    assert page.text == "振込番号  TRF96500401"


# --------------------------------------------------------------------------- office files
def test_docx_body_tables_headers_and_footers():
    import docx

    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "Iberia Componentes S.L."
    d.add_paragraph("Aviso de pago")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Factura", "Importe"
    t.cell(1, 0).text, t.cell(1, 1).text = "FA2026-09035", "4.310,00"
    d.add_paragraph("Saludos")
    d.sections[0].footer.paragraphs[0].text = "Barcelona"
    buf = io.BytesIO()
    d.save(buf)
    r = parse_document("aviso.docx", None, buf.getvalue())
    assert r.method == "docx"
    assert r.text.splitlines()[0] == "Iberia Componentes S.L." and r.text.splitlines()[-1] == "Barcelona"
    assert r.text.index("Aviso de pago") < r.text.index("FA2026-09035") < r.text.index("Saludos")
    assert [t["name"] for t in r.tables] == ["body"]
    assert r.tables[0]["rows"] == [["Factura", "Importe"], ["FA2026-09035", "4.310,00"]]


def test_xlsx_values_and_iso_dates():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Payment"
    ws.append(["Invoice", "Date", "Amount"])
    ws.append(["JP-26-002500", dt.datetime(2026, 8, 28), 473000.0])
    ws.append(["JP-26-002464", dt.date(2026, 8, 19), 1606824.5])
    buf = io.BytesIO()
    wb.save(buf)
    r = parse_document("plan.xlsx", None, buf.getvalue())
    assert r.method == "spreadsheet" and r.tables[0]["name"] == "Payment"
    assert r.tables[0]["rows"][1:] == [["JP-26-002500", "2026-08-28", "473000"], ["JP-26-002464", "2026-08-19", "1606824.5"]]


def test_csv_semicolon():
    r = parse_document("pay.csv", None, "Rechnung;Betrag\n91005442;15.999,58\n".encode("cp1252"))
    assert r.method == "csv" and r.tables[0]["rows"] == [["Rechnung", "Betrag"], ["91005442", "15.999,58"]]


def test_html_attachment_and_size_limit(monkeypatch):
    html = b"<html><body><p>Paid:</p><table><tr><th>Invoice</th><th>Amount</th></tr>" \
           b"<tr><td>UK-1043514</td><td>GBP 1,707.60</td></tr></table></body></html>"
    r = parse_document("advice.html", None, html)
    assert r.tables[0]["rows"] == [["Invoice", "Amount"], ["UK-1043514", "GBP 1,707.60"]]
    assert "UK-1043514\tGBP 1,707.60" in r.text
    monkeypatch.setattr(get_settings(), "html_max_bytes", 10)
    assert parse_document("advice.html", None, html).status == "SKIPPED"


def test_containers_and_unknown_types_are_skipped():
    assert parse_document("PaymentNotice.zip", None, b"PK\x03\x04...").method == "container"
    r = parse_document("mail.msg", None, b"\xd0\xcf\x11\xe0")
    assert (r.method, r.status) == ("unsupported", "SKIPPED")
    assert kind_of("x.TIF", None, b"II*\x00") == "image"


# --------------------------------------------------------------------------- bodies
def test_plain_text_body_table_r22():
    body = load_sample("R22-0128").body
    layout, lang = parse_body(body)
    rows = layout["tables"][0]["rows"]
    assert rows[0] == ["Invoice No.", "Amount paid"]
    assert rows[1:] == [["Inv# UK-1043514", "GBP 1,707.60"], ["UK-1043671", "GBP 1,673.51"], ["UK-1043497", "GBP 3,223.69"]]
    assert lang == "en"


def test_plain_text_body_table_keeps_empty_amount_r02():
    rows = parse_body(load_sample("R02-0019").body)[0]["tables"][0]["rows"]
    assert rows == [["Rechnungs-Nr.", "Zahlbetrag"], ["91005442", ""]]


def test_html_body_tables():
    html = "<p>Dear AR,</p><table><tr><td>Invoice</td><td>Amount</td></tr><tr><td>90012345</td><td>EUR 1.234,56</td></tr></table>"
    layout, _ = parse_body(html_to_text(html), html)
    assert layout["format"] == "html" and layout["tables"][0]["rows"][1] == ["90012345", "EUR 1.234,56"]
    assert find_html_tables("<table><tr><td>a<br>b</td></tr></table>") == [[["a b"]]]


def test_prose_is_not_a_table():
    assert find_text_tables("Dear Sir,\n\nKindly note our payment as per the attached advice.\n") == []


def test_language_detection_is_repeatable():
    text = "Wir haben folgende Zahlung veranlasst. Mit freundlichen Gruessen, Elbe Chemie GmbH"
    assert {detect_language(text) for _ in range(5)} == {"de"}


# --------------------------------------------------------------------------- database: parse_email, reuse
def _hide_existing_copies(db, email):
    """The laptop database may already hold a parsed copy of the same sample file (which would be
    reused). Mark those PENDING inside the test's transaction (rolled back afterwards)."""
    shas = db.scalars(select(EmailAttachment.sha256).where(EmailAttachment.email_id == email.id)).all()
    for a in db.scalars(select(EmailAttachment).where(EmailAttachment.sha256.in_(shas),
                                                      EmailAttachment.email_id != email.id)):
        a.parse_status = "PENDING"
    db.flush()


def test_parse_email_saves_results_and_reuses_identical_files(db, storage):
    raw = build_message(load_sample("R22-0128")).as_bytes()  # body-only
    e1 = store_email(db, raw, storage=storage, gmail_id="m02-a")
    msg = build_message(load_sample("R01-0007"))
    e2 = store_email(db, msg.as_bytes(), storage=storage, gmail_id="m02-b")
    e3 = store_email(db, msg.as_bytes(), storage=storage, gmail_id="m02-c")  # same PNG again
    db.commit()
    _hide_existing_copies(db, e2)

    s1 = parse_email(db, e1, storage=storage)
    assert (s1.parsed, s1.body_tables, e1.status, e1.language) == (0, 1, "PARSED", "en")
    assert e1.body_layout["tables"][0]["rows"][1] == ["Inv# UK-1043514", "GBP 1,707.60"]

    s2 = parse_email(db, e2, storage=storage)
    [a2] = db.scalars(select(EmailAttachment).where(EmailAttachment.email_id == e2.id)).all()
    assert (s2.parsed, a2.parse_status, a2.parse_method, a2.parse_reused) == (1, "PARSED", "ocr", False)
    assert "UK-1043360" in a2.text_content and a2.layout["pages"][0]["words"] and a2.parse_ms > 0

    s3 = parse_email(db, e3, storage=storage)
    [a3] = db.scalars(select(EmailAttachment).where(EmailAttachment.email_id == e3.id)).all()
    assert (s3.parsed, s3.reused, a3.parse_reused) == (0, 1, True)
    assert a3.text_content == a2.text_content and a3.parse_ms == a2.parse_ms and s3.saved_ms == a2.parse_ms


def test_parse_failure_is_recorded_not_raised(db, storage, monkeypatch):
    e = store_email(db, build_message(load_sample("R01-0007")).as_bytes(), storage=storage, gmail_id="m02-t")
    db.commit()
    monkeypatch.setattr(get_settings(), "doc_timeout_s", 0.001)
    _hide_existing_copies(db, e)  # otherwise the PNG would be reused instead of read (and timed out)
    s = parse_email(db, e, storage=storage)
    [a] = db.scalars(select(EmailAttachment).where(EmailAttachment.email_id == e.id)).all()
    assert (s.failed, a.parse_status, e.status) == (1, "FAILED", "FAILED") and a.parse_error.startswith("timeout")


def test_admin_parse_endpoint_needs_key(monkeypatch):
    monkeypatch.setattr(get_settings(), "admin_api_key", "k" * 20)
    c = TestClient(app)
    assert c.post("/api/admin/emails/EML-999999/parse").status_code == 401
    assert c.post("/api/admin/emails/EML-999999/parse", headers={"X-Admin-Key": "k" * 20}).status_code == 404
